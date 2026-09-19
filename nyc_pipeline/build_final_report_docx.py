from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "final_report.md"
OUTPUT = ROOT / "Shanghai_New_York_Park_Opening_Comparative_Technical_Report_2026_08_07.docx"

TABLE_CAPTIONS = [
    "Validated New York focal park-opening events",
    "Central continuous-exposure PPML estimates, lambda = 500",
    "New York binary-threshold PPML estimates, 500 m",
    "Joint pretrend tests for event times -3 and -2",
    "New York continuous-exposure design diagnostics, lambda = 500",
]


def set_run_font(run, size=12, bold=None, italic=None, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill="FFFFFF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, values in edges.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in values.items():
            element.set(qn("w:" + key), str(value))


def set_three_line_table(table):
    table.autofit = True
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rows = table.rows
    if not rows:
        return
    clear = {"val": "nil"}
    top = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    mid = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
    bottom = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    for r_idx, row in enumerate(rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell)
            set_cell_border(cell, top=clear, bottom=clear, left=clear, right=clear, insideH=clear, insideV=clear)
            if r_idx == 0:
                set_cell_border(cell, top=top, bottom=mid, left=clear, right=clear)
            if r_idx == len(rows) - 1:
                set_cell_border(cell, bottom=bottom, left=clear, right=clear)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, 9.5, bold=(r_idx == 0))
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, 9)


def add_inline(paragraph, text, size=12):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size, font="Courier New")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size)


def add_equation(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_rpr = OxmlElement("m:rPr")
    math_style = OxmlElement("m:sty")
    math_style.set(qn("m:val"), "p")
    math_rpr.append(math_style)
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_rpr)
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    paragraph._p.append(math_para)


def parse_table(lines, start):
    table_lines = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        table_lines.append(lines[idx].strip())
        idx += 1
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        rows.pop(1)
    return rows, idx


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for style_name, size, before, after in [
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 13, 10, 4),
        ("Heading 3", 12, 8, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_count = 0
    figure_number = 0
    table_number = 0
    idx = 0
    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()
        if not line:
            idx += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(72)
            p.paragraph_format.space_after = Pt(18)
            run = p.add_run(line[2:])
            set_run_font(run, 18, bold=True)
            title_count += 1
            idx += 1
            continue

        if line.startswith("## ") and title_count == 1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            run = p.add_run(line[3:])
            set_run_font(run, 14, bold=True)
            title_count += 1
            idx += 1
            continue

        if title_count == 2 and line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:-2])
            set_run_font(run, 11, italic=True)
            p.add_run().add_break(WD_BREAK.PAGE)
            title_count += 1
            idx += 1
            continue

        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
            idx += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            idx += 1
            continue

        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", line)
        if image_match:
            figure_number += 1
            image_path = ROOT / image_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run()
            run.add_picture(str(image_path), width=Inches(5.65))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Cm(0)
            caption.paragraph_format.line_spacing = 1.0
            caption.paragraph_format.space_after = Pt(6)
            r = caption.add_run(f"Figure {figure_number}. {image_match.group(1)}.")
            set_run_font(r, 10.5, italic=True)
            idx += 1
            continue

        if line.startswith("|"):
            rows, next_idx = parse_table(lines, idx)
            table_number += 1
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.first_line_indent = Cm(0)
            cap.paragraph_format.line_spacing = 1.0
            cap.paragraph_format.keep_with_next = True
            caption_text = TABLE_CAPTIONS[table_number - 1] if table_number <= len(TABLE_CAPTIONS) else "Results"
            r = cap.add_run(f"Table {table_number}. {caption_text}")
            set_run_font(r, 10.5, italic=True)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r_idx, row_data in enumerate(rows):
                for c_idx, value in enumerate(row_data):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = value
            set_three_line_table(table)
            note = doc.add_paragraph()
            note.paragraph_format.first_line_indent = Cm(0)
            note.paragraph_format.line_spacing = 1.0
            note.paragraph_format.space_after = Pt(5)
            text = "Notes: Standard errors are in parentheses where reported. Statistical significance is evaluated at p < 0.05; no 10 percent stars are used."
            if table_number == 1:
                text = "Notes: Areas use the current official NYC Parks property polygon. The strict-opening robustness sample removes replacement and rebuilt-opening events."
            elif table_number == 4:
                text = "Notes: The joint null is that the event-time -3 and -2 exposure coefficients both equal zero. Event time -1 is the omitted reference period."
            elif table_number == 5:
                text = "Notes: Entries are PPML log coefficients with standard errors in parentheses. The strict sample contains six first/new public openings. Overlap exclusion removes cells within 1,500 m of another focal park already open at the event date."
            rr = note.add_run(text)
            set_run_font(rr, 9.5, italic=True)
            idx = next_idx
            continue

        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            idx += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_inline(p, line[2:])
            idx += 1
            continue

        if line.startswith("**") and line.endswith("**") and any(
            token in line for token in ("K_up", "A_upt", "E[Y_upt")
        ):
            equation = line[2:-2]
            equation = equation.replace("K_up", "Kᵤₚ").replace("A_upt", "Aᵤₚₜ")
            equation = equation.replace("d_up", "dᵤₚ").replace("Post_pt", "Postₚₜ")
            equation = equation.replace("Y_upt", "Yᵤₚₜ").replace("alpha_up", "αᵤₚ")
            equation = equation.replace("delta_pt", "δₚₜ").replace("beta", "β").replace("lambda", "λ")
            equation = equation.replace(" x ", " × ")
            add_equation(doc, equation)
            idx += 1
            continue

        paragraph_lines = [line]
        idx += 1
        while idx < len(lines):
            nxt = lines[idx].strip()
            if not nxt:
                idx += 1
                break
            if nxt.startswith(("#", "|", "- ", "![")) or re.match(r"^\d+\. ", nxt):
                break
            paragraph_lines.append(nxt)
            idx += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(paragraph_lines))

    props = doc.core_properties
    props.title = "Park Openings and Recorded Crime in Shanghai and New York City"
    props.subject = "Reproducible cross-city event-study replication"
    props.author = "Hongyang Zhang"
    props.keywords = "parks; crime; Shanghai; New York City; PPML; event study"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
