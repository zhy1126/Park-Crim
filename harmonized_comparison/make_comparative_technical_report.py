from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
RESULTS = ROOT / "results"
ARTIFACTS = ROOT / "artifacts"
FIGURES = ROOT / "figures"
REPORTS = ROOT / "reports"
REPORTS.mkdir(parents=True, exist_ok=True)
OUTPUT = REPORTS / "harmonized_shanghai_nyc_comparative_technical_report_2026_08_09_revised.docx"

BLACK = RGBColor(0, 0, 0)
CITY_LABELS = {"SH": "Shanghai", "NYC": "New York City"}
OUTCOME_LABELS = {
    "total_crime": "Total crime",
    "theft": "Theft / larceny",
    "non_theft": "Non-theft",
    "day_crime": "Daytime crime",
    "night_crime": "Nighttime crime",
}
OUTCOMES = list(OUTCOME_LABELS)


def read_csv(name: str) -> pd.DataFrame:
    return pd.read_csv(RESULTS / name)


SHORT_STATIC = read_csv("city_specific_static_ppml.csv")
LONG_STATIC = read_csv("city_specific_static_ppml_long_m3_p3.csv")
SHORT_DIFF = read_csv("pooled_city_difference_tests.csv")
LONG_DIFF = read_csv("pooled_city_difference_tests_long_m3_p3.csv")
SHORT_PRE = read_csv("joint_pretrend_tests.csv")
LONG_PRE = read_csv("joint_pretrend_tests_long_m3_p3.csv")
SHORT_PLACEBO = read_csv("preperiod_placebo_tests.csv")
LONG_PLACEBO = read_csv("preperiod_placebo_tests_long_m3_p3.csv")
SHORT_LOO = read_csv("leave_one_event_out.csv")
LONG_LOO = read_csv("leave_one_event_out_long_m3_p3.csv")
SHORT_JACK = read_csv("park_jackknife_inference.csv")
LONG_JACK = read_csv("park_jackknife_inference_long_m3_p3.csv")
SUPPORT = read_csv("outcome_support_diagnostics.csv")
LEDGER = pd.read_csv(DATA / "harmonized_event_ledger.csv")
SHORT_PANEL = pd.read_csv(DATA / "harmonized_stacked_grid_panel_500m.csv", low_memory=False)
LONG_PANEL = pd.read_csv(DATA / "harmonized_stacked_grid_panel_500m_long_m3_p3.csv", low_memory=False)


def set_font(run, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)
    section.left_margin = Mm(31.8)
    section.right_margin = Mm(31.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    heading_specs = {
        "Heading 1": (14, True, False, 10, 4),
        "Heading 2": (12, True, True, 8, 3),
        "Heading 3": (12, True, False, 6, 2),
    }
    for style_name, (size, bold, italic, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = BLACK
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.0


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_font(run, 10)


def add_body(doc: Document, text: str, indent: bool = True) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.first_line_indent = Cm(0.74 if indent else 0)
    run = paragraph.add_run(text)
    set_font(run, 12)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.37)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_font(run, 12)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    set_font(run, 14 if level == 1 else 12, bold=True, italic=(level == 2))


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_cell_border(cell, edge: str, val: str = "single", size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    tag = qn(f"w:{edge}")
    border = tc_borders.find(tag)
    if border is None:
        border = OxmlElement(f"w:{edge}")
        tc_borders.append(border)
    border.set(qn("w:val"), val)
    if val != "nil":
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:color"), "000000")


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_cm: list[float]) -> None:
    widths_dxa = [int(round(width_cm * 567)) for width_cm in widths_cm]
    total_dxa = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "90")
    tbl_ind.set(qn("w:type"), "dxa")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_dxa in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width_dxa))
        grid.append(column)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width_dxa = widths_dxa[idx]
            cell.width = Cm(widths_cm[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")


def add_table_caption(doc: Document, caption: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(caption)
    set_font(run, 10, italic=True)


def add_three_line_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths_cm: list[float],
    note: str | None = None,
    alignments: list[int] | None = None,
) -> None:
    add_table_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_cm)
    set_repeat_header(table.rows[0])
    alignments = alignments or [WD_ALIGN_PARAGRAPH.CENTER] * len(headers)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignments[idx]
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(header)
        set_font(run, 9.5, bold=True)
        set_cell_border(cell, "top", "single", 10)
        set_cell_border(cell, "bottom", "single", 6)
        set_cell_border(cell, "left", "nil")
        set_cell_border(cell, "right", "nil")

    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignments[idx]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.first_line_indent = Cm(0)
            run = paragraph.add_run(str(value))
            set_font(run, 9.5)
            set_cell_border(cell, "top", "nil")
            set_cell_border(cell, "bottom", "nil")
            set_cell_border(cell, "left", "nil")
            set_cell_border(cell, "right", "nil")

    for cell in table.rows[-1].cells:
        set_cell_border(cell, "bottom", "single", 10)

    if note:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(4)
        lead = paragraph.add_run("Notes: ")
        set_font(lead, 9, italic=True)
        run = paragraph.add_run(note)
        set_font(run, 9)


def add_figure(doc: Document, filename: str, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(FIGURES / filename), width=Cm(14.35))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", filename)
    caption_p = doc.add_paragraph(style="Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.first_line_indent = Cm(0)
    caption_p.paragraph_format.space_before = Pt(2)
    caption_p.paragraph_format.space_after = Pt(5)
    caption_run = caption_p.add_run(caption)
    set_font(caption_run, 10, italic=True)


def star(p_value: float) -> str:
    if pd.isna(p_value):
        return ""
    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"
    return ""


def coefficient(row: pd.Series) -> str:
    return f"{row['estimate']:.3f}{star(float(row['p_value']))} ({row['std_error']:.3f})"


def static_row(data: pd.DataFrame, city: str, outcome: str, design: str, sample: str = "core") -> pd.Series:
    rows = data[
        (data["city_code"] == city)
        & (data["outcome"] == outcome)
        & (data["design"] == design)
        & (data["sample"] == sample)
        & (data["vcov"] == "grid")
    ]
    if len(rows) != 1:
        raise ValueError(f"Expected one row: {city}, {outcome}, {design}, {sample}; got {len(rows)}")
    return rows.iloc[0]


def difference_row(data: pd.DataFrame, outcome: str, design: str, sample: str = "core") -> pd.Series:
    rows = data[
        (data["outcome"] == outcome)
        & (data["design"] == design)
        & (data["sample"] == sample)
        & (data["vcov"] == "grid")
    ]
    if len(rows) != 1:
        raise ValueError(f"Expected one difference row: {outcome}, {design}, {sample}; got {len(rows)}")
    return rows.iloc[0]


def contrast(row: pd.Series) -> str:
    return f"{row['estimate_nyc_minus_sh']:.3f}{star(float(row['p_value']))} ({row['std_error']:.3f})"


def sample_summary(panel: pd.DataFrame, long: bool = False) -> list[list[str]]:
    if long:
        panel = panel[panel["analysis_core"] == 1].copy()
    rows = []
    for city in ["SH", "NYC"]:
        data = panel[panel["city_code"] == city]
        rows.append(
            [
                CITY_LABELS[city],
                f"{data['event_id'].nunique():,}",
                f"{data['grid_cluster_id'].nunique():,}",
                f"{data['stack_unit_id'].nunique():,}",
                f"{len(data):,}",
            ]
        )
    return rows


def build_report() -> Path:
    doc = Document()
    configure_document(doc)
    add_page_number(doc.sections[0])
    core_properties = doc.core_properties
    core_properties.title = "Harmonized Shanghai-New York Park-Opening Comparative Study"
    core_properties.subject = "Technical report and publication-readiness assessment"
    core_properties.author = "Research workflow output"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Harmonized Park-Opening Exposure and Recorded Cases in Shanghai and New York City")
    set_font(run, 16, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run("Technical Report and Publication-Readiness Assessment")
    set_font(run, 12, italic=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.first_line_indent = Cm(0)
    meta.paragraph_format.space_after = Pt(10)
    run = meta.add_run("Frozen analysis build: 9 August 2026")
    set_font(run, 10)

    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "This report reconstructs the Shanghai and New York City park-opening analyses under one spatial, temporal, and econometric protocol. The harmonized design contains 11 researcher-verified Shanghai events and 9 New York City events supported by official municipal records. Both cities use 500 m square grids, a 1,500 m event support, six-month periods centered on each opening date, distance from the park boundary, and Poisson pseudo-maximum-likelihood models with stack-by-grid and stack-by-period fixed effects.",
    )
    add_body(
        doc,
        "The all-event short window covers the three six-month periods before opening and the first six months after opening. Shanghai adjudicated judgment records exhibit a positive total-case gradient under continuous distance-decayed exposure (beta = 0.566, SE = 0.194, p = 0.004), whereas the corresponding New York City complaint-record estimate is smaller and imprecise (beta = 0.059, SE = 0.033, p = 0.080). The direct continuous city contrast is detectable (NYC minus Shanghai = -0.507, p = 0.010), but it is not reproduced by the binary 500 m design (p = 0.136). The Shanghai continuous coefficient implies approximately 10.8 percent higher expected recorded counts for a one-standard-deviation exposure increase; the much larger full-scale contrast from exposure zero to one is not a realistic average policy effect.",
    )
    add_body(
        doc,
        "A longer balanced [-3,+3] window is available for only 6 Shanghai and 9 New York City events. In that sample, the average Shanghai total-case coefficient weakens to 0.318 (p = 0.114), while the nighttime coefficient remains positive (0.703, p = 0.006). Re-estimating the short model on the same 15 long-eligible events yields a Shanghai total coefficient of 0.575 (p = 0.026). The attenuation in the long average therefore appears to reflect the post-opening time path more than event-composition change alone. Dynamic coefficients are non-monotonic, however, so they do not establish a stable permanent effect.",
    )
    add_body(
        doc,
        "The correct publication frame is a measurement-aware comparative replication. The results compare within-system proportional changes in Shanghai adjudicated judgments and New York City police complaints; they do not compare equal levels of underlying crime or equivalent stages of criminal-justice processing. The build is ready as a harmonized data-and-methods report after artifact synchronization, but it is not ready to support a causal cross-city effect claim without outcome-stage harmonization and stronger few-event inference.",
    )

    add_heading(doc, "1. Research Questions and Comparative Estimand", 1)
    add_body(
        doc,
        "The first research question asks whether recorded cases change more strongly after a focal park opens in grid cells that are closer to that park. The second asks whether this post-opening spatial gradient differs between the two city-specific administrative record systems. These questions concern differential exposure gradients within focal-event stacks. They do not ask whether citywide crime increased, nor do they identify a city-level difference in latent offending incidence.",
    )
    add_bullet(doc, "RQ1: Within each city, do more-exposed grid cells exhibit different post-opening changes in recorded cases than less-exposed cells in the same park-opening stack?")
    add_bullet(doc, "RQ2: Are the estimated within-system exposure gradients similar in Shanghai adjudicated judgments and New York City police complaints?")
    add_bullet(doc, "RQ3: How sensitive are the gradients to exposure definition, follow-up horizon, event composition, cross-park overlap, outcome sparsity, and omission of individual parks?")

    add_heading(doc, "2. Harmonized Data Construction", 1)
    add_heading(doc, "2.1 Event Ledger", 2)
    add_body(
        doc,
        "The Shanghai ledger contains all 11 opening dates supplied and manually verified by the researcher. These dates are treated as project-verified inputs and are not independently re-adjudicated by the comparative workflow. No Shanghai event is excluded from the ledger or the all-event short-window analysis. The balanced long-window estimator uses six Shanghai events because five verified events lack complete follow-up through event time +3. The New York City ledger contains 9 openings or substantial reconstruction reopenings documented by official NYC Parks records. Each event is linked to a polygon boundary and retains date precision, area, event class, and provenance metadata. Appendix Tables A1 and A2 list the complete city-specific event ledgers.",
    )
    add_heading(doc, "2.2 Spatial and Temporal Support", 2)
    add_body(
        doc,
        "For every park, the workflow creates a separate stack of 500 m square grid cells whose polygons intersect a 1,500 m support around the park boundary. Distance is measured from each grid polygon to the focal park polygon, so cells intersecting the park receive zero boundary distance. Event time is centered on the recorded opening date rather than on calendar half-years. Period -1 is the six months immediately before opening, period 0 is the first six months after opening, and periods -3 through +3 define the longer dynamic window.",
    )
    add_body(
        doc,
        "The short window retains all 20 events and contains 3,836 stacked observations. The long panel contains 6,713 constructed rows, but only 4,977 rows from 15 events have complete [-3,+3] support and enter the balanced long models. The five later Shanghai openings remain valid events; they are absent from the long estimator only because the source period does not provide complete future follow-up.",
    )

    add_body(doc, "Table 1 summarizes the common spatial support and makes the short-versus-long event composition explicit.")
    table_rows = []
    for window, panel in [("Short [-3,0]", SHORT_PANEL), ("Long [-3,+3]", LONG_PANEL[LONG_PANEL["analysis_core"] == 1])]:
        for city in ["SH", "NYC"]:
            subset = panel[panel["city_code"] == city]
            table_rows.append(
                [
                    window,
                    CITY_LABELS[city],
                    f"{subset['event_id'].nunique():,}",
                    f"{subset['grid_cluster_id'].nunique():,}",
                    f"{subset['stack_unit_id'].nunique():,}",
                    f"{len(subset):,}",
                ]
            )
    add_three_line_table(
        doc,
        "Table 1. Harmonized panel support by city and analysis window",
        ["Window", "City", "Events", "Unique grids", "Stack-units", "Rows"],
        table_rows,
        [2.4, 3.2, 1.4, 2.1, 2.1, 2.0],
        note="The long-window rows include only events with complete event times -3 through +3. Grid identifiers are city-specific; stack-units are park-by-grid units.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "2.3 Outcomes and Source-Stage Difference", 2)
    add_body(
        doc,
        "The estimated outcomes are total recorded cases, theft or larceny, non-theft, daytime cases, and nighttime cases. Daytime is defined as 06:00-17:59 and nighttime as 18:00-05:59. An unknown-time count is retained only as an accounting field to verify that total equals day plus night plus unknown; it is not estimated as a separate outcome. Shanghai observations are adjudicated judgment records, whereas New York City observations are police complaint records. This source-stage difference is the principal limit on substantive cross-city equivalence.",
    )
    add_figure(
        doc,
        "figure_3_outcome_support_comparison.png",
        "Figure 1. Outcome support in the harmonized short-window panels. Shanghai contains a much larger share of always-zero stack-units, while New York City contains a much larger share of nonzero complaint observations. The contrast reflects both case density and administrative-record stage.",
    )
    add_body(
        doc,
        "For total crime, 55.7 percent of Shanghai stack-units are always zero, compared with 14.8 percent in New York City. For Shanghai non-theft, the always-zero share rises to 82.7 percent. PPML is appropriate for nonnegative counts and retains zero observations where they contribute within-unit variation, but all-zero fixed-effect groups do not identify a coefficient and drop from the effective likelihood. Main-table observation counts must therefore be interpreted as outcome-specific active samples.",
    )

    add_heading(doc, "3. Identification and Estimation", 1)
    add_heading(doc, "3.1 Continuous Exposure", 2)
    add_body(
        doc,
        "The main exposure is the product of a time-invariant negative-exponential distance weight and an indicator for periods after the focal park opens. With decay parameter 500 m, a grid touching the park has weight 1, a grid approximately 347 m away has weight about 0.5, and more distant grids approach zero. Each stack uses distance to its own focal park; exposure is not summed across all parks.",
    )
    add_body(
        doc,
        "The conditional mean is estimated by PPML with stack-by-grid fixed effects and stack-by-event-period fixed effects. The former absorb all stable differences between a grid and its focal park, including the baseline distance weight. The latter absorb shocks common to all grids in the same park stack and period, including the stack-wide post-opening shift. The coefficient is therefore identified by whether grids with higher pre-determined distance weights change differently after opening than less-exposed grids within the same stack. This is a continuous-treatment difference-in-differences design and requires counterfactual trends to be comparable across exposure intensity.",
    )

    add_heading(doc, "3.2 Binary Exposure and Pooled City Contrast", 2)
    add_body(
        doc,
        "The binary robustness model replaces the smooth distance weight with an indicator for grid cells within 500 m of the focal park after opening. This specification tests whether the main pattern survives without relying on the exponential dose-response shape. Pooled models interact treatment with city indicators and report the NYC-minus-Shanghai coefficient difference. Because the two outcomes occupy different administrative stages, this interaction is interpreted as a difference in recorded-case gradients, not a structural difference in underlying offending.",
    )

    add_heading(doc, "3.3 Dynamic and Diagnostic Specifications", 2)
    add_body(
        doc,
        "Dynamic models replace the static post indicator with exposure-by-event-time interactions and omit event time -1. Joint Wald tests assess whether the -3 and -2 coefficients equal zero. Pre-period placebo models assign false openings at -2 and -1. Cross-park overlap is audited separately: 16.5 percent of stacked units are ever exposed to another focal opening during the short window, and an uncontaminated sample excludes them. Leave-one-event-out estimates and park jackknife summaries assess dependence on individual focal parks. Grid-clustered standard errors are primary; ordinary park-clustered estimates are not treated as decisive because there are only 11 and 9 event clusters, falling to 6 and 9 in the long sample.",
    )

    add_heading(doc, "4. Short-Window Results", 1)
    add_body(
        doc,
        "Figure 2 and Table 2 report the all-event short-window estimates. The continuous Shanghai total-case coefficient is positive and statistically distinguishable from zero under conventional grid-clustered inference, although precision remains constrained by only 11 park events and substantial outcome sparsity. New York City coefficients are smaller across outcomes. The binary specification preserves the direction of the Shanghai estimates but weakens precision. Accordingly, the evidence is strongest for a short-run continuous exposure gradient and does not establish an exposure-definition-independent binary treatment effect.",
    )
    add_figure(
        doc,
        "figure_1_static_city_comparison.png",
        "Figure 2. City-specific short-window PPML estimates. The sample contains all 11 Shanghai and 9 New York City events. Points are coefficients and whiskers are 95 percent confidence intervals based on grid-clustered standard errors.",
    )

    rows = []
    for outcome in OUTCOMES:
        rows.append(
            [
                OUTCOME_LABELS[outcome],
                coefficient(static_row(SHORT_STATIC, "SH", outcome, "continuous_park_opening_ppml")),
                coefficient(static_row(SHORT_STATIC, "NYC", outcome, "continuous_park_opening_ppml")),
                coefficient(static_row(SHORT_STATIC, "SH", outcome, "binary_park_opening_ppml")),
                coefficient(static_row(SHORT_STATIC, "NYC", outcome, "binary_park_opening_ppml")),
            ]
        )
    add_three_line_table(
        doc,
        "Table 2. Short-window city-specific PPML estimates",
        ["Outcome", "SH continuous", "NYC continuous", "SH binary", "NYC binary"],
        rows,
        [3.0, 2.8, 2.8, 2.8, 2.8],
        note="Entries are PPML coefficients with grid-clustered standard errors in parentheses. Continuous exposure uses lambda = 500 m; binary exposure equals one within 500 m after opening. * p < 0.05, ** p < 0.01, *** p < 0.001.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )

    contrast_rows = []
    for outcome in OUTCOMES:
        continuous = difference_row(SHORT_DIFF, outcome, "continuous_park_opening_ppml")
        binary = difference_row(SHORT_DIFF, outcome, "binary_park_opening_ppml")
        contrast_rows.append(
            [
                OUTCOME_LABELS[outcome],
                contrast(continuous),
                f"{continuous['p_value']:.4f}",
                contrast(binary),
                f"{binary['p_value']:.4f}",
            ]
        )
    add_body(
        doc,
        "Table 3 tests the city differences directly, avoiding the invalid inference that two city estimates differ merely because one is statistically significant and the other is not.",
    )
    add_three_line_table(
        doc,
        "Table 3. Direct short-window city-difference tests",
        ["Outcome", "Continuous NYC-SH", "p", "Binary NYC-SH", "p"],
        contrast_rows,
        [3.0, 3.6, 1.4, 3.6, 1.4],
        note="Negative estimates indicate a smaller New York City gradient than the Shanghai gradient. The coefficient equality test is meaningful only as a comparison between recorded-case systems.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_body(
        doc,
        "The statistically detectable continuous total-case difference is not reproduced by the binary contrast. Shanghai theft is also imprecise, and non-theft is particularly sparse. The most responsible summary is therefore that Shanghai displays a larger short-run continuous recorded-case gradient, with nighttime cases contributing a notable secondary pattern, while New York City estimates are small and statistically imprecise. It would be incorrect to infer a robust city difference from the fact that one city-specific coefficient is significant and the other is not.",
    )

    add_heading(doc, "5. Horizon, Event Composition, and Dynamic Results", 1)
    add_body(
        doc,
        "A direct short-versus-long comparison changes both follow-up horizon and Shanghai event composition unless the event set is held fixed. Table 4 therefore reports three quantities: the short estimate using all events, the short estimate restricted to the 15 events with complete long support, and the long average using those same 15 events. For Shanghai total cases, the short coefficient is nearly unchanged when the sample is restricted from 11 to 6 events, but the longer average weakens. The attenuation is therefore not explained solely by event composition and may reflect follow-up horizon, dynamic heterogeneity, and changes in the outcome-specific active PPML sample.",
    )

    horizon_rows = []
    for city, outcome in [("SH", "total_crime"), ("SH", "night_crime"), ("NYC", "total_crime"), ("NYC", "night_crime")]:
        short_all = static_row(SHORT_STATIC, city, outcome, "continuous_park_opening_ppml", "core")
        short_fixed = static_row(SHORT_STATIC, city, outcome, "continuous_park_opening_ppml", "long_eligible_events")
        long_average = static_row(LONG_STATIC, city, outcome, "continuous_park_opening_ppml", "core")
        horizon_rows.append(
            [
                CITY_LABELS[city],
                OUTCOME_LABELS[outcome],
                coefficient(short_all),
                coefficient(short_fixed),
                coefficient(long_average),
            ]
        )
    add_three_line_table(
        doc,
        "Table 4. Separating event composition from follow-up horizon",
        ["City", "Outcome", "Short: all", "Short: long-eligible", "Long: same events"],
        horizon_rows,
        [2.6, 2.7, 3.0, 3.0, 3.0],
        note="Continuous exposure estimates are shown. The long-eligible set contains 6 Shanghai and 9 New York City events. The short restricted model retains event times -3 through 0; the long model uses -3 through +3.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 3,
    )

    add_figure(
        doc,
        "figure_2_long_dynamic_event_study.png",
        "Figure 3. Long-window continuous-exposure event-study estimates. The balanced sample contains 6 Shanghai and 9 New York City events. Event time -1 is the omitted reference period; whiskers are 95 percent confidence intervals based on grid-clustered standard errors.",
    )
    add_body(
        doc,
        "Shanghai total-case coefficients are positive at opening and at +3 but are negative at +1 and +2; none provides a clean monotonic persistence pattern. Shanghai nighttime coefficients are positive at opening, +2, and +3, but the sequence includes a negative +1 coefficient and only six events. The +3 nighttime estimate is statistically detectable in the pointwise model, yet it arises among many outcome-by-period comparisons and lacks simultaneous confidence bands. It is best treated as a secondary, hypothesis-generating temporal pattern. New York City dynamic estimates remain close to zero and imprecise.",
    )

    add_heading(doc, "6. Pretrends, Placebos, and Few-Event Sensitivity", 1)
    add_body(
        doc,
        "Table 5 summarizes the joint pre-opening lead tests and within-window placebo checks for total and nighttime recorded cases in both analysis windows.",
    )
    diag_rows = []
    for city in ["SH", "NYC"]:
        for outcome in ["total_crime", "night_crime"]:
            short_pre = SHORT_PRE[
                (SHORT_PRE["city_code"] == city)
                & (SHORT_PRE["outcome"] == outcome)
                & (SHORT_PRE["treatment_type"] == "continuous")
            ].iloc[0]
            long_pre = LONG_PRE[
                (LONG_PRE["city_code"] == city)
                & (LONG_PRE["outcome"] == outcome)
                & (LONG_PRE["treatment_type"] == "continuous")
            ].iloc[0]
            short_placebo = SHORT_PLACEBO[
                (SHORT_PLACEBO["city_code"] == city)
                & (SHORT_PLACEBO["outcome"] == outcome)
                & SHORT_PLACEBO["design"].str.startswith("continuous_")
            ]["p_value"].min()
            long_placebo = LONG_PLACEBO[
                (LONG_PLACEBO["city_code"] == city)
                & (LONG_PLACEBO["outcome"] == outcome)
                & LONG_PLACEBO["design"].str.startswith("continuous_")
            ]["p_value"].min()
            diag_rows.append(
                [
                    CITY_LABELS[city],
                    OUTCOME_LABELS[outcome],
                    f"{short_pre['p_value']:.4f}",
                    f"{short_placebo:.4f}",
                    f"{long_pre['p_value']:.4f}",
                    f"{long_placebo:.4f}",
                ]
            )
    add_three_line_table(
        doc,
        "Table 5. Continuous-exposure timing diagnostics",
        ["City", "Outcome", "Short joint lead p", "Short min placebo p", "Long joint lead p", "Long min placebo p"],
        diag_rows,
        [2.6, 2.6, 2.3, 2.3, 2.3, 2.3],
        note="Joint lead tests evaluate event times -3 and -2. The placebo column reports the smaller p-value across false openings at -2 and -1. Failure to reject is not proof of parallel trends, especially with few events and sparse outcomes.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    add_body(
        doc,
        "No reported joint lead or placebo diagnostic crosses the 5 percent threshold, but several are close. The smallest short placebo p-value is 0.0504 for New York City total complaints, and the smallest long joint lead p-value is approximately 0.0597 for New York City daytime complaints. The NYC short and long placebo entries reuse the same pre-period records and are not independent validations. These checks do not show strong evidence against parallel trends in the selected joint tests, but they neither confirm the identifying assumption nor rule out anticipatory divergence. Their power is limited by the number of events and by outcome-specific PPML attrition.",
    )

    add_body(
        doc,
        "Figure 4 then examines whether the continuous-exposure estimates are disproportionately driven by any single focal park.",
    )
    add_figure(
        doc,
        "figure_4_leave_one_event_out_stability.png",
        "Figure 4. Leave-one-event-out continuous-exposure estimates. The short window contains 11 Shanghai and 9 New York City events; the long window contains 6 and 9. Each point omits one focal park, and horizontal bars mark the mean across omissions. Sign stability is a dependence diagnostic and does not replace valid small-cluster inference.",
    )
    add_body(
        doc,
        "All Shanghai leave-one-event-out continuous total-case estimates remain positive in both windows. Short-window values range from 0.511 to 0.638, while long-window values remain positive but smaller. Shanghai nighttime estimates are also positive under each omission. Park-jackknife summaries yield small p-values for selected outcomes, but these should be read as stability diagnostics rather than definitive Type I error control, particularly when the long Shanghai sample contains only six event clusters.",
    )

    jack_rows = []
    for window, source in [("Short", SHORT_JACK), ("Long", LONG_JACK)]:
        for city, outcome in [("SH", "total_crime"), ("SH", "night_crime"), ("NYC", "total_crime"), ("NYC", "night_crime")]:
            row = source[
                (source["city_code"] == city)
                & (source["outcome"] == outcome)
                & (source["design"] == "continuous_park_opening_ppml")
            ].iloc[0]
            jack_rows.append(
                [
                    window,
                    CITY_LABELS[city],
                    OUTCOME_LABELS[outcome],
                    f"{row['leave_one_out_min']:.3f} to {row['leave_one_out_max']:.3f}",
                    f"{row['same_sign_share']:.2f}",
                    "<0.001" if row["park_jackknife_p"] < 0.001 else f"{row['park_jackknife_p']:.3f}",
                ]
            )
    add_body(
        doc,
        "Table 6 condenses these leave-one-event-out ranges and park-level jackknife summaries for the principal total and nighttime outcomes.",
    )
    add_three_line_table(
        doc,
        "Table 6. Leave-one-event-out and park-jackknife diagnostics",
        ["Window", "City", "Outcome", "LOO range", "Same-sign share", "Jackknife p"],
        jack_rows,
        [1.7, 2.5, 2.5, 3.0, 2.3, 2.0],
        note="LOO denotes leave-one-event-out. The jackknife treats focal parks as the resampling clusters. With 6-11 clusters, the results summarize event dependence and should not be treated as a complete small-sample inferential correction.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 3,
    )

    add_heading(doc, "7. Additional Robustness Boundaries", 1)
    add_heading(doc, "7.1 Cross-Park Exposure", 2)
    add_body(
        doc,
        "The stacked exposure is defined relative to one focal park, but some grid-periods can be near another park that opens during the same window. The audit identifies 16.5 percent of stack-units as ever subject to such competing exposure. The positive Shanghai direction is not eliminated by excluding flagged overlapping units. Because this restriction changes the comparison sample and cannot remove all co-development, it does not establish that cross-park exposure is irrelevant.",
    )
    add_heading(doc, "7.2 Spatial and Outcome Sparsity", 2)
    add_body(
        doc,
        "Changing spatial support alters the number of all-zero units and the effective PPML sample. The common 500 m grid is retained because it provides the same nominal support in both cities, but this does not make the recorded outcomes equally dense. The Shanghai non-theft result is especially fragile because only 5.4 percent of observations are nonzero. Cross-outcome rankings should therefore be avoided: each PPML outcome is estimated on its own active fixed-effect sample.",
    )
    add_heading(doc, "7.3 Event Type", 2)
    add_body(
        doc,
        "The event mix is not identical. Shanghai is dominated by substantial reopenings and rebuilds, whereas New York City includes more first openings. A narrow rebuild-comparable sample is available, but it contains only 6 Shanghai and 3 New York City events. It is informative as a directional sensitivity check but too small to support a separate definitive city comparison.",
    )

    add_heading(doc, "8. Publication-Readiness Assessment", 1)
    readiness_rows = [
        ["Event provenance", "Achieved", "11 Shanghai events are researcher-verified; 9 NYC events have official municipal provenance."],
        ["Spatial and temporal harmonization", "Achieved", "Common 500 m grids, 1,500 m support, boundary distance, and exact-date-centered six-month periods."],
        ["Reproducible estimation", "Achieved", "All short and long PPML models converged; identities and balance checks pass."],
        ["Outcome-stage equivalence", "Unresolved", "Shanghai judgments and NYC complaints represent different criminal-justice selection processes."],
        ["Few-event inference", "Partial", "Grid clustering, leave-one-out, and jackknife diagnostics are available, but only 6-11 park clusters remain."],
        ["Total-case robustness", "Partial", "Short continuous contrast is detectable; binary and long-average contrasts are not."],
        ["Dynamic interpretation", "Exploratory", "Shanghai paths are non-monotonic and selected late coefficients face multiple-testing concerns."],
    ]
    add_body(
        doc,
        "Table 7 distinguishes the elements that are fully reproducible from the remaining limitations that prevent a strong structural cross-city causal claim.",
    )
    add_three_line_table(
        doc,
        "Table 7. Publication-readiness matrix",
        ["Criterion", "Status", "Assessment"],
        readiness_rows,
        [3.5, 2.1, 8.7],
        note="Status refers to the current frozen comparative build, not to the validity of the original city-specific datasets for other research questions.",
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "8.1 Defensible Claim", 2)
    add_body(
        doc,
        "Under the continuous distance-decay specification, the common design produces a larger short-run recorded-case gradient for Shanghai adjudicated judgments than for New York City complaints. The corresponding binary 500 m city contrast is not statistically distinguishable from zero, and the total-case contrast also weakens in the longer balanced sample. These are specification-dependent associations in different administrative-record systems, not evidence that parks affect underlying crime differently across cities.",
    )
    add_heading(doc, "8.2 Claims to Avoid", 2)
    add_bullet(doc, "Do not state that park openings caused a 76 percent increase in crime. Report the approximately 10.8 percent one-standard-deviation short Shanghai contrast and identify the outcome as adjudicated recorded cases.")
    add_bullet(doc, "Do not state that there is no effect in New York City. State that the NYC estimates are small and statistically imprecise under the present sample and design.")
    add_bullet(doc, "Do not state that parallel trends are confirmed. State that the joint lead tests do not reject, while power is limited and some diagnostics are marginal.")
    add_bullet(doc, "Do not state that the Shanghai effect persists through two years after opening. Event time +3 covers months 18-24 after opening; the dynamic path is non-monotonic and the long Shanghai sample contains only six events.")
    add_bullet(doc, "Do not treat smaller park-clustered or two-way standard errors as stronger evidence. Conventional asymptotics are unreliable with so few park clusters.")

    add_heading(doc, "8.3 Highest-Value Next Steps", 2)
    add_bullet(doc, "Harmonize the administrative outcome stage if feasible, either by obtaining police-record outcomes for Shanghai or prosecution/adjudication outcomes for New York City.")
    add_bullet(doc, "Pre-specify one primary outcome, exposure, and time window before journal submission; keep nighttime and offense-type patterns explicitly secondary.")
    add_bullet(doc, "Add a nonlinear few-cluster inference procedure with documented finite-sample behavior, or present randomization-style inference tied to a defensible event assignment mechanism.")
    add_bullet(doc, "Use simultaneous confidence bands for the dynamic event study and distinguish pointwise exploratory coefficients from family-wise inference.")
    add_bullet(doc, "Expand the event ledger, especially the number of long-follow-up Shanghai events, without relaxing the researcher-verified timing rule.")

    add_heading(doc, "9. Reproducibility Statement", 1)
    add_body(
        doc,
        "All comparative files are contained in the harmonized_comparison project directory. Existing Shanghai and New York City source data are read without deletion, renaming, movement, or overwrite. The event ledger, short and long stacked panels, model outputs, figures, plotting data, diagnostic JSON files, code, artifact manifest, and evaluation file are stored as separate local artifacts. SHA-256 hashes are recorded in artifacts\\artifact_manifest.csv.",
    )
    add_bullet(doc, "Authoritative event ledger: data\\harmonized_event_ledger.csv")
    add_bullet(doc, "Short panel: data\\harmonized_stacked_grid_panel_500m.csv")
    add_bullet(doc, "Long panel: data\\harmonized_stacked_grid_panel_500m_long_m3_p3.csv")
    add_bullet(doc, "Model code: run_harmonized_models.R")
    add_bullet(doc, "Figure code: make_comparative_figures.py")
    add_bullet(doc, "Mechanical evaluator: evaluate.py")
    add_bullet(doc, "Artifact status note: ARTIFACT_STATUS.md")

    doc.add_page_break()
    add_heading(doc, "Appendix A. Harmonized Event Ledger", 1)
    for city_code, caption in [("SH", "Table A1. Researcher-verified Shanghai focal events"), ("NYC", "Table A2. Official-source New York City focal events")]:
        rows = []
        subset = LEDGER[LEDGER["city_code"] == city_code].sort_values("opening_date")
        for _, item in subset.iterrows():
            rows.append(
                [
                    str(item["park_name"]),
                    str(item["opening_date"]),
                    str(item["date_precision"]),
                    str(item["event_family"]).replace("_", " "),
                    f"{float(item['boundary_area_ha']):.2f}",
                ]
            )
        add_three_line_table(
            doc,
            caption,
            ["Park", "Opening date", "Precision", "Event class", "Area (ha)"],
            rows,
            [4.2, 2.4, 1.7, 4.2, 1.6],
            note=(
                "All Shanghai dates are researcher-verified project inputs and no Shanghai event is excluded."
                if city_code == "SH"
                else "NYC dates and event descriptions are supported by official NYC Parks records retained in the ledger."
            ),
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
        )

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_report()
    print(json.dumps({"report": str(output), "generated": date.today().isoformat()}, indent=2))
